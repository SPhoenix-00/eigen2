"""
Transform trading CSV into daily position tracking table.

Reads the committee trading results and creates a multi-sheet Excel workbook:
- Positions: number of shares held (floor of coefficient)
- Prices: Google Finance formulas for historical prices
- Buy Flag: 1 on days the agent buys, 0 otherwise
- Sell Flag: 1 on days the agent sells, 0 otherwise
- Summary: Portfolio value, Peak Capital, and Cumulative P&L
"""

import pandas as pd
import math
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


def parse_date(date_str: str) -> datetime:
    """Parse date string in various formats."""
    # Handle various date formats
    formats = [
        "%d-%m-%y",      # 02-01-24
        "%Y-%m-%d",      # 2024-01-02
        "%d/%m/%y",      # 02/01/24
        "%d/%m/%Y",      # 02/01/2024
        "%Y-%m-%d %H:%M:%S",  # 2024-01-02 00:00:00
    ]

    for fmt in formats:
        try:
            return datetime.strptime(str(date_str).split()[0] if ' ' in str(date_str) else str(date_str), fmt)
        except ValueError:
            continue

    raise ValueError(f"Cannot parse date: {date_str}")


def extract_date_range_from_filename(filename: str) -> tuple:
    """
    Extract start and end dates from filename pattern like:
    '..._29-07-22_to_10-03-23.csv' (with hyphens)
    or
    '..._290722_to_100323.csv' (without hyphens)

    Returns:
        (start_date, end_date) as datetime objects, or (None, None) if not found
    """
    import re

    # Pattern 1: DD-MM-YY_to_DD-MM-YY (with hyphens)
    pattern1 = r'(\d{2}-\d{2}-\d{2})_to_(\d{2}-\d{2}-\d{2})'
    match = re.search(pattern1, filename)
    if match:
        start_str, end_str = match.groups()
        return parse_date(start_str), parse_date(end_str)

    # Pattern 2: DDMMYY_to_DDMMYY (without hyphens)
    pattern2 = r'(\d{6})_to_(\d{6})'
    match = re.search(pattern2, filename)
    if match:
        start_str, end_str = match.groups()
        # Convert DDMMYY to DD-MM-YY
        start_formatted = f"{start_str[:2]}-{start_str[2:4]}-{start_str[4:]}"
        end_formatted = f"{end_str[:2]}-{end_str[2:4]}-{end_str[4:]}"
        return parse_date(start_formatted), parse_date(end_formatted)

    return None, None


def transform_trades_to_positions(csv_path: str, output_path: str = None,
                                   start_date: datetime = None,
                                   end_date: datetime = None) -> pd.DataFrame:
    """
    Transform trading CSV to multi-sheet Excel workbook.

    Args:
        csv_path: Path to trades CSV file
        output_path: Optional path for output Excel file
        start_date: Optional explicit start date for the date range
        end_date: Optional explicit end date for the date range

    If start_date/end_date are not provided, attempts to extract from filename,
    then falls back to the actual trade dates in the CSV.
    """
    # Read the trading data
    df = pd.read_csv(csv_path)

    # Parse dates
    df['entry_date_parsed'] = df['entry_date'].apply(parse_date)
    df['exit_date_parsed'] = df['exit_date'].apply(parse_date)

    # Determine date range (priority: explicit params > filename > trade dates)
    if start_date is not None and end_date is not None:
        min_date, max_date = start_date, end_date
    else:
        # Try to get from filename
        input_path = Path(csv_path)
        min_date, max_date = extract_date_range_from_filename(input_path.name)

        if min_date is None or max_date is None:
            # Fall back to trade dates
            min_date = df['entry_date_parsed'].min()
            max_date = df['exit_date_parsed'].max()

    # Create all calendar days in the range
    all_dates = pd.date_range(start=min_date, end=max_date, freq='D')
    num_days = len(all_dates)

    # Get unique stocks (sorted)
    stocks = sorted(df['stock_name'].unique())
    num_stocks = len(stocks)

    # Initialize DataFrames
    positions = pd.DataFrame(0, index=all_dates, columns=stocks)
    buy_flags = pd.DataFrame(0, index=all_dates, columns=stocks)
    sell_flags = pd.DataFrame(0, index=all_dates, columns=stocks)

    # Fill in data from trades
    for _, trade in df.iterrows():
        stock = trade['stock_name']
        entry = trade['entry_date_parsed']
        exit_date = trade['exit_date_parsed']
        coefficient = trade.get('coefficient', 1.0)
        quantity = int(math.floor(coefficient))

        # Positions: held from entry up to (not including) exit
        mask = (positions.index >= entry) & (positions.index < exit_date)
        positions.loc[mask, stock] += quantity

        # Buy flag: 1 on entry date
        if entry in buy_flags.index:
            buy_flags.loc[entry, stock] = 1

        # Sell flag: 1 on exit date
        if exit_date in sell_flags.index:
            sell_flags.loc[exit_date, stock] = 1

    # Format index as date strings
    date_strings = all_dates.strftime('%Y-%m-%d').tolist()
    positions.index = date_strings
    buy_flags.index = date_strings
    sell_flags.index = date_strings

    # Determine output path
    if output_path is None:
        output_path = input_path.parent / f"{input_path.stem}_positions.xlsx"

    # Create workbook
    wb = Workbook()

    # --- Sheet 1: Positions ---
    ws_pos = wb.active
    ws_pos.title = "Positions"
    _write_df_to_sheet(ws_pos, positions)

    # --- Sheet 2: Prices (Google Finance formulas) ---
    ws_prices = wb.create_sheet("Prices")
    # Header row: date column + stock tickers
    ws_prices.cell(row=1, column=1, value="date")
    for col_idx, stock in enumerate(stocks, start=2):
        ws_prices.cell(row=1, column=col_idx, value=stock)

    # Data rows with GOOGLEFINANCE formulas (with IFERROR fallback to previous day)
    for row_idx, date_str in enumerate(date_strings, start=2):
        ws_prices.cell(row=row_idx, column=1, value=date_str)
        for col_idx, stock in enumerate(stocks, start=2):
            col_letter = get_column_letter(col_idx)
            # =IFERROR(INDEX(GOOGLEFINANCE(B$1,"price",$A2),2,2), B1)
            gf_formula = f'INDEX(GOOGLEFINANCE({col_letter}$1,"price",$A{row_idx}),2,2)'
            fallback = f'{col_letter}{row_idx - 1}'
            formula = f'=IFERROR({gf_formula},{fallback})'
            ws_prices.cell(row=row_idx, column=col_idx, value=formula)

    # --- Sheet 3: Buy Flag ---
    ws_buy = wb.create_sheet("Buy Flag")
    _write_df_to_sheet(ws_buy, buy_flags)

    # --- Sheet 4: Sell Flag ---
    ws_sell = wb.create_sheet("Sell Flag")
    _write_df_to_sheet(ws_sell, sell_flags)

    # --- Sheet 5: Summary ---
    ws_summary = wb.create_sheet("Summary")
    ws_summary.cell(row=1, column=1, value="date")
    ws_summary.cell(row=1, column=2, value="Stock Value")
    ws_summary.cell(row=1, column=3, value="Cash")
    ws_summary.cell(row=1, column=4, value="Portfolio Value")
    ws_summary.cell(row=1, column=5, value="Daily Buy Cost")
    ws_summary.cell(row=1, column=6, value="Peak Capital")
    ws_summary.cell(row=1, column=7, value="Daily Sell Revenue")
    ws_summary.cell(row=1, column=8, value="Cumulative P&L")

    # Build formulas for each row
    # Column references for other sheets
    first_stock_col = "B"
    last_stock_col = get_column_letter(num_stocks + 1)
    stock_range = f"{first_stock_col}:{last_stock_col}"

    for row_idx, date_str in enumerate(date_strings, start=2):
        ws_summary.cell(row=row_idx, column=1, value=date_str)

        # Column B: Stock Value = SUMPRODUCT(Positions row * Prices row)
        stock_value_formula = f"=SUMPRODUCT(Positions!{first_stock_col}{row_idx}:{last_stock_col}{row_idx},Prices!{first_stock_col}{row_idx}:{last_stock_col}{row_idx})"
        ws_summary.cell(row=row_idx, column=2, value=stock_value_formula)

        # Column C: Cash = cumulative sell revenue (no outflow for buys)
        if row_idx == 2:
            cash_formula = "=G2"
        else:
            cash_formula = f"=C{row_idx - 1}+G{row_idx}"
        ws_summary.cell(row=row_idx, column=3, value=cash_formula)

        # Column D: Portfolio Value = Stock Value + Cash
        portfolio_formula = f"=B{row_idx}+C{row_idx}"
        ws_summary.cell(row=row_idx, column=4, value=portfolio_formula)

        # Column E: Daily Buy Cost = SUMPRODUCT(Buy Flag row * Positions row * Prices row)
        buy_cost_formula = f"=SUMPRODUCT('Buy Flag'!{first_stock_col}{row_idx}:{last_stock_col}{row_idx},Positions!{first_stock_col}{row_idx}:{last_stock_col}{row_idx},Prices!{first_stock_col}{row_idx}:{last_stock_col}{row_idx})"
        ws_summary.cell(row=row_idx, column=5, value=buy_cost_formula)

        # Column F: Peak Capital = cumulative sum of Daily Buy Cost
        if row_idx == 2:
            peak_capital_formula = "=E2"
        else:
            peak_capital_formula = f"=F{row_idx - 1}+E{row_idx}"
        ws_summary.cell(row=row_idx, column=6, value=peak_capital_formula)

        # Column G: Daily Sell Revenue = SUMPRODUCT(Sell Flag row * Positions[prev day] * Prices row)
        prev_row = row_idx - 1 if row_idx > 2 else row_idx
        sell_revenue_formula = f"=SUMPRODUCT('Sell Flag'!{first_stock_col}{row_idx}:{last_stock_col}{row_idx},Positions!{first_stock_col}{prev_row}:{last_stock_col}{prev_row},Prices!{first_stock_col}{row_idx}:{last_stock_col}{row_idx})"
        ws_summary.cell(row=row_idx, column=7, value=sell_revenue_formula)

        # Column H: Cumulative P&L = Portfolio Value - Peak Capital
        pnl_formula = f"=D{row_idx}-F{row_idx}"
        ws_summary.cell(row=row_idx, column=8, value=pnl_formula)

    # Save workbook
    wb.save(output_path)

    print(f"Saved position tracking workbook to: {output_path}")
    print(f"Date range: {min_date.strftime('%Y-%m-%d')} to {max_date.strftime('%Y-%m-%d')}")
    print(f"Total days: {num_days}")
    print(f"Total stocks: {num_stocks}")
    print(f"Sheets: Positions, Prices, Buy Flag, Sell Flag, Summary")

    return positions


def _write_df_to_sheet(ws, df):
    """Write a DataFrame to an openpyxl worksheet."""
    # Header row
    ws.cell(row=1, column=1, value="date")
    for col_idx, col_name in enumerate(df.columns, start=2):
        ws.cell(row=1, column=col_idx, value=col_name)

    # Data rows
    for row_idx, (date_str, row_data) in enumerate(df.iterrows(), start=2):
        ws.cell(row=row_idx, column=1, value=date_str)
        for col_idx, value in enumerate(row_data, start=2):
            ws.cell(row=row_idx, column=col_idx, value=value)


def main():
    results_dir = Path(__file__).parent / "committee_results"
    csv_files = list(results_dir.glob("*.csv"))
    csv_files = [f for f in csv_files if "_positions" not in f.name]

    if not csv_files:
        print("No trading CSV files found in committee_results/")
        return

    for csv_file in csv_files:
        print(f"\nProcessing: {csv_file.name}")
        print("-" * 50)
        positions = transform_trades_to_positions(str(csv_file))

        print(f"\nPosition summary (first 5 days):")
        print(positions.head())


if __name__ == "__main__":
    main()
